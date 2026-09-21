"""DOCX export built natively from the resume JSON (spec §13 requires this
NOT be derived from the rendered PDF) using python-docx, so the output is a
true, selectable, ATS-parseable Word document.
"""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Pt

from app.pdf.link_builder import build_personal_links
from app.pdf.renderer import RenderSettings
from app.schemas.resume import ResumeData

LINK_COLOR = "2563EB"  # matches the PDF template's brand-color link styling


def add_hyperlink(paragraph, url: str, text: str):
    """python-docx has no built-in hyperlink support; this builds the
    required OOXML by hand so the exported DOCX has a real, clickable
    link (not underlined, colored to match the PDF template) rather than
    plain unclickable text.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    run_props.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    run_props.append(underline)

    run.append(run_props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _add_contact_line(doc, items: list[dict]):
    """items: [{"text": str, "href": str | None}, ...] rendered centered,
    separated by " | ", with href items as real clickable hyperlinks."""
    if not items:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, item in enumerate(items):
        if i > 0:
            p.add_run(" | ")
        if item["href"]:
            add_hyperlink(p, item["href"], item["text"])
        else:
            p.add_run(item["text"])


def render_resume_docx(resume: ResumeData, settings: RenderSettings) -> bytes:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = settings.font_family
    style.font.size = Pt(settings.font_size_pt)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = _inches(settings.margin_in)
        section.left_margin = section.right_margin = _inches(settings.margin_in)

    info = resume.personal_info

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name_p.add_run(info.name)
    run.bold = True
    run.font.size = Pt(settings.font_size_pt + 9)

    if info.title:
        p = doc.add_paragraph(info.title)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    links = build_personal_links(info.model_dump(by_alias=True))

    line1 = []
    if info.location:
        line1.append({"text": info.location, "href": None})
    if info.email:
        line1.append({"text": info.email, "href": links["email"]})
    if info.phone:
        line1.append({"text": info.phone, "href": None})
    if info.linkedin:
        line1.append({"text": info.linkedin, "href": links["linkedin"]})
    _add_contact_line(doc, line1)

    line2 = []
    if info.github:
        line2.append({"text": f"GitHub: {info.github}", "href": links["github"]})
    if info.portfolio:
        line2.append({"text": f"Portfolio: {info.portfolio}", "href": links["portfolio"]})
    if info.leetcode:
        line2.append({"text": f"LeetCode: {info.leetcode}", "href": links["leetcode"]})
    if info.hackerrank:
        line2.append({"text": f"HackerRank: {info.hackerrank}", "href": links["hackerrank"]})
    _add_contact_line(doc, line2)

    for section in resume.sections:
        if not section.visible:
            continue
        heading = doc.add_heading(section.title, level=2)
        for run in heading.runs:
            run.font.size = Pt(settings.heading_size_pt)

        content = section.content
        if section.type == "summary":
            doc.add_paragraph(content.get("text", ""))

        elif section.type == "skills":
            for cat in content.get("categories", []):
                p = doc.add_paragraph()
                p.add_run(f"{cat['name']}: ").bold = True
                p.add_run(", ".join(cat.get("items", [])))

        elif section.type == "experience":
            for item in content.get("items", []):
                p = doc.add_paragraph()
                header = item.get("title", "")
                if item.get("company"):
                    header += f" | {item['company']}"
                dates = " – ".join(v for v in [item.get("startDate", ""), item.get("endDate", "")] if v)
                p.add_run(header).bold = True
                if dates:
                    p.add_run(f"\t{dates}")
                for bullet in item.get("bullets", []):
                    doc.add_paragraph(bullet, style="List Bullet")

        elif section.type == "projects":
            for item in content.get("items", []):
                p = doc.add_paragraph()
                p.add_run(item.get("name", "")).bold = True
                if item.get("technologies"):
                    doc.add_paragraph(", ".join(item["technologies"]))
                for bullet in item.get("bullets", []):
                    doc.add_paragraph(bullet, style="List Bullet")

        elif section.type == "education":
            for item in content.get("items", []):
                line = item.get("degree", "")
                if item.get("institution"):
                    line += f" | {item['institution']}"
                extras = " | ".join(v for v in [item.get("gpa", ""), item.get("year", "")] if v)
                if extras:
                    line += f" | {extras}"
                doc.add_paragraph(line)

        elif section.type == "certifications":
            for item in content.get("items", []):
                parts = [item.get("name", "")]
                if item.get("issuer"):
                    parts.append(item["issuer"])
                if item.get("year"):
                    parts.append(item["year"])
                if item.get("credentialUrl"):
                    parts.append(f"Credential: {item['credentialUrl']}")
                doc.add_paragraph(" | ".join(p for p in parts if p))

        else:
            if content.get("text"):
                doc.add_paragraph(content["text"])
            for i in content.get("items", []) or []:
                if isinstance(i, str):
                    doc.add_paragraph(i, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _inches(value: float):
    from docx.shared import Inches

    return Inches(value)
