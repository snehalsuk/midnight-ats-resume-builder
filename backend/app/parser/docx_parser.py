import io

import docx


def extract_lines_from_docx(file_bytes: bytes) -> list[str]:
    document = docx.Document(io.BytesIO(file_bytes))
    lines: list[str] = []
    for para in document.paragraphs:
        stripped = para.text.strip()
        if stripped:
            lines.append(stripped)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                stripped = cell.text.strip()
                if stripped:
                    lines.append(stripped)
    return lines
